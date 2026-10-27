"""Импорт быстрых ответов (Excel) и полезных ссылок (txt) в library.quick_answers / links.

Разбор файлов (parse_*) не знает о базе; запись (import_*) получает уже открытую сессию и сама
ничего не коммитит: транзакцией владеет вызывающий (эндпоинт открывает соединение, и при любой
ошибке не остаётся ни одной записанной строки).

Эмбеддинги здесь не считаются. Строки ложатся с embedding = NULL, триггеры миграции 011 будят
ingest-worker, а тот встраивает quick_answer_variants.text и links.description КАК ХРАНИТСЯ.
Поэтому варианты сохраняются уже нормализованными: по тому же правилу их нормализует
graph-rag-svc при поиске.
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import load_workbook

COL_QUESTION = "Основной вопрос"
COL_ANSWER = "Ответ на русском"
COL_VARIANTS = "Варианты вопроса на русском"

# Правило ноутбука, воспроизведённое дословно (общий контракт с graph-rag-svc). Диапазон
# [а-яА-Я] не включает Ё/ё, поэтому буква вырезается: «ещё» превращается в «ещ».
# Править нельзя: запрос и сохранённые варианты должны нормализоваться одинаково.
_STRIP_RE = re.compile(r"[^a-zA-Zа-яА-Я0-9\s]")
_SPLIT_RE = re.compile(r"[\r\n;]+")
_URL_RE = re.compile(r"https?://")


class ImportFileError(ValueError):
    """Файл не годится для импорта (формат, кодировка, нет нужных колонок)."""


def normalize(text: str) -> str:
    return _STRIP_RE.sub("", text).lower().strip()


def _dedup_normalized(texts: list[str]) -> list[str]:
    seen: dict[str, None] = {}
    for t in texts:
        n = normalize(t)
        if n:
            seen.setdefault(n, None)
    return list(seen)


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_quick_answers_xlsx(data: bytes) -> list[dict]:
    """Строки листа -> [{"question", "answer", "variants"}]. Колонки ищутся по заголовкам.

    Полностью пустые строки отбрасываются здесь; строки с пустым вопросом или ответом
    возвращаются как есть, их отсеет и посчитает импортёр (skipped).
    """
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # zipfile.BadZipFile, KeyError и пр.: любой нечитаемый файл
        raise ImportFileError("Не удалось прочитать xlsx") from exc
    try:
        ws = wb["Ответы"] if "Ответы" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ImportFileError("Лист пустой")
        names = [_cell(h) for h in header]
        try:
            iq, ia, iv = (names.index(c) for c in (COL_QUESTION, COL_ANSWER, COL_VARIANTS))
        except ValueError:
            # Колонка вариантов необязательна: без неё остаётся основной вопрос.
            if COL_QUESTION in names and COL_ANSWER in names:
                iq, ia, iv = names.index(COL_QUESTION), names.index(COL_ANSWER), None
            else:
                raise ImportFileError(
                    f"Нет колонок «{COL_QUESTION}» и «{COL_ANSWER}» в первой строке листа"
                ) from None
        out: list[dict] = []
        for row in rows:
            def at(i):
                return _cell(row[i]) if i is not None and i < len(row) else ""
            question, answer, raw_variants = at(iq), at(ia), at(iv)
            if not (question or answer or raw_variants):
                continue
            parts = ([question] if question else []) + _SPLIT_RE.split(raw_variants)
            out.append({"question": question, "answer": answer, "variants": _dedup_normalized(parts)})
        return out
    finally:
        wb.close()


def parse_links_txt(text: str) -> list[dict]:
    """Строки «описание. URL» -> [{"description", "url"}]; строки без URL пропускаются."""
    out: list[dict] = []
    for line in text.splitlines():
        m = _URL_RE.search(line)
        if not m:
            continue
        description = line[: m.start()].strip().rstrip(". ").strip()
        out.append({"description": description, "url": line[m.start():].strip()})
    return out


# ---------- запись в БД ----------
# Все запросы параметризованы; сессия — psycopg.AsyncConnection (или подделка в тестах).

SQL_QA_SELECT = "SELECT id, question, answer FROM library.quick_answers WHERE question = ANY(%s)"
SQL_VARIANT_SELECT = (
    "SELECT quick_answer_id, text FROM library.quick_answer_variants WHERE quick_answer_id = ANY(%s)"
)
SQL_QA_INSERT = "INSERT INTO library.quick_answers (question, answer) VALUES (%s, %s) RETURNING id"
SQL_QA_UPDATE_ANSWER = "UPDATE library.quick_answers SET answer = %s, updated_at = now() WHERE id = %s"
SQL_QA_TOUCH = "UPDATE library.quick_answers SET updated_at = now() WHERE id = %s"
SQL_VARIANT_RESET = (
    "UPDATE library.quick_answer_variants SET embedding = NULL, embedding_model = NULL "
    "WHERE quick_answer_id = %s"
)
SQL_VARIANT_INSERT = (
    "INSERT INTO library.quick_answer_variants (quick_answer_id, text) VALUES (%s, %s) "
    "ON CONFLICT (quick_answer_id, text) DO NOTHING"
)
SQL_VARIANT_DELETE = (
    "DELETE FROM library.quick_answer_variants WHERE quick_answer_id = %s AND text = ANY(%s)"
)
SQL_LINK_SELECT = "SELECT url, description FROM library.links WHERE url = ANY(%s)"
SQL_LINK_INSERT = "INSERT INTO library.links (description, url) VALUES (%s, %s)"
SQL_LINK_UPDATE = (
    "UPDATE library.links SET description = %s, embedding = NULL, embedding_model = NULL WHERE url = %s"
)


async def import_quick_answers(session, rows: list[dict]) -> dict:
    """Синхронизирует вопросы из файла: upsert по question.

    Для каждого вопроса из файла лишние варианты удаляются, новые добавляются, прежние остаются
    вместе с эмбеддингами. Изменился answer — эмбеддинги всех его вариантов сбрасываются, и воркер
    считает их заново. Вопросы, которых в файле нет, не трогаются. updated_at меняется только
    при реальной правке. Повторный импорт того же файла ничего не пишет.
    """
    created = updated = skipped = 0
    wanted: dict[str, dict] = {}
    for row in rows:
        question = _cell(row.get("question"))
        answer = _cell(row.get("answer"))
        variants = _dedup_normalized([question] + list(row.get("variants") or []))
        # Вопрос без ответа или целиком из знаков препинания, а также повтор вопроса в файле.
        if not question or not answer or not normalize(question) or question in wanted:
            skipped += 1
            continue
        wanted[question] = {"answer": answer, "variants": variants}
    if not wanted:
        return {"created": 0, "updated": 0, "skipped": skipped}

    cur = await session.execute(SQL_QA_SELECT, (list(wanted),))
    existing = {q: (i, a) for i, q, a in await cur.fetchall()}
    cur = await session.execute(SQL_VARIANT_SELECT, ([i for i, _ in existing.values()],))
    have: dict[Any, set[str]] = {}
    for qa_id, text in await cur.fetchall():
        have.setdefault(qa_id, set()).add(text)

    for question, item in wanted.items():
        target = item["variants"]
        if question not in existing:
            cur = await session.execute(SQL_QA_INSERT, (question, item["answer"]))
            (qa_id,) = await cur.fetchone()
            for text in target:
                await session.execute(SQL_VARIANT_INSERT, (qa_id, text))
            created += 1
            continue
        qa_id, old_answer = existing[question]
        current = have.get(qa_id, set())
        add = [t for t in target if t not in current]
        drop = sorted(current - set(target))
        answer_changed = old_answer != item["answer"]
        if not (answer_changed or add or drop):
            continue
        if answer_changed:
            await session.execute(SQL_QA_UPDATE_ANSWER, (item["answer"], qa_id))
            await session.execute(SQL_VARIANT_RESET, (qa_id,))
        else:
            await session.execute(SQL_QA_TOUCH, (qa_id,))
        if drop:
            await session.execute(SQL_VARIANT_DELETE, (qa_id, drop))
        for text in add:
            await session.execute(SQL_VARIANT_INSERT, (qa_id, text))
        updated += 1
    return {"created": created, "updated": updated, "skipped": skipped}


async def import_links(session, rows: list[dict]) -> dict:
    """Upsert ссылок по url. Описание изменилось — эмбеддинг сбрасывается; ссылки, которых нет в
    файле, остаются как есть."""
    created = updated = skipped = 0
    wanted: dict[str, str] = {}
    for row in rows:
        description = _cell(row.get("description"))
        url = _cell(row.get("url"))
        if not description or not url or url in wanted:
            skipped += 1
            continue
        wanted[url] = description
    if not wanted:
        return {"created": 0, "updated": 0, "skipped": skipped}

    cur = await session.execute(SQL_LINK_SELECT, (list(wanted),))
    existing = dict(await cur.fetchall())
    for url, description in wanted.items():
        if url not in existing:
            await session.execute(SQL_LINK_INSERT, (description, url))
            created += 1
        elif existing[url] != description:
            await session.execute(SQL_LINK_UPDATE, (description, url))
            updated += 1
    return {"created": created, "updated": updated, "skipped": skipped}
